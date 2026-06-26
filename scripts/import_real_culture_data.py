from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.database import Database
from app.services.text_utils import normalize_text
from tagger.preprocess import preprocess_lyrics


BASE_DIR = Path(__file__).resolve().parent.parent
RAW_DIR = BASE_DIR / "data" / "raw"
PROCESSED_DIR = BASE_DIR / "data" / "processed"

MELON_CHART_PATH = RAW_DIR / "music" / "melon_monthly_top50.csv"
MELON_LYRICS_PATH = RAW_DIR / "music" / "melon_song_lyrics.csv"
KYOBO_PATH = (
    RAW_DIR
    / "교보문고_연도별_베스트셀러(2000~2025)"
    / "kyobo_annual_bestsellers_2000_2025.csv"
)
KOBIS_SEED_PATH = RAW_DIR / "movie" / "seed.sql"

TOKEN_RE = re.compile(r"[가-힣A-Za-z0-9]{2,}")

STOPWORDS = {
    "그리고",
    "그래서",
    "하지만",
    "나는",
    "너는",
    "우리는",
    "그대",
    "우리",
    "너를",
    "나를",
    "다시",
    "아직",
    "정말",
    "오늘",
    "내일",
    "지금",
    "그냥",
    "모든",
    "없는",
    "있는",
    "하게",
    "해서",
    "보다",
    "처럼",
}

EMOTION_RULES: dict[str, list[str]] = {
    "슬픔": ["이별", "눈물", "슬프", "그립", "외로", "아프", "떠나", "울", "어둠", "상처"],
    "그리움": ["그리", "기억", "추억", "다시", "돌아", "기다", "보고 싶"],
    "분노": ["분노", "미워", "화가", "억울", "배신", "소리쳐", "참을 수"],
    "불안": ["불안", "두려", "걱정", "떨", "흔들", "막막", "길을 잃"],
    "무기력": ["지쳐", "피곤", "버거", "힘들", "멈춰", "무너", "끝난"],
    "기쁨": ["행복", "웃", "좋아", "신나", "축하", "반짝", "빛나"],
    "설렘": ["설레", "두근", "처음", "고백", "기대", "만나"],
    "감사": ["고마", "감사", "덕분", "소중"],
    "평온": ["평온", "잔잔", "조용", "쉬어", "바람", "하늘", "별", "안녕"],
    "희망": ["희망", "꿈", "내일", "빛", "살아", "괜찮", "일어나", "앞으로"],
    "호기심": ["궁금", "찾아", "모험", "새로운", "세계", "비밀"],
}

TOPIC_RULES: dict[str, list[str]] = {
    "사랑": ["사랑", "마음", "연인", "고백", "너와", "함께"],
    "이별": ["이별", "떠나", "헤어", "그리움", "눈물"],
    "위로": ["위로", "괜찮", "안아", "쉬어", "치유"],
    "성장": ["성장", "도전", "꿈", "성공", "미래", "시작"],
    "가족": ["가족", "엄마", "아빠", "부모", "아이"],
    "사회": ["사회", "민중", "자유", "세상", "역사", "정의"],
    "일상": ["하루", "집", "길", "밤", "아침", "시간"],
    "여행": ["여행", "바다", "기차", "길", "도시"],
    "판타지": ["마법", "모험", "비밀", "세계", "왕국"],
    "경제": ["돈", "부자", "투자", "경제", "기업", "재테크"],
    "인문": ["철학", "역사", "인간", "사회", "문명", "삶"],
    "과학": ["과학", "우주", "기술", "생명", "실험"],
    "학습": ["공부", "학습", "시험", "지식", "교훈"],
}


def clean(value: Any) -> str:
    return normalize_text(str(value or ""))


def norm_key(*parts: Any) -> str:
    joined = " ".join(clean(part).lower() for part in parts)
    return re.sub(r"[^0-9a-z가-힣]+", "", joined)


def stable_id(prefix: str, key: str) -> str:
    digest = hashlib.blake2b(key.encode("utf-8"), digest_size=8).hexdigest()
    return f"{prefix}_{digest}"


def parse_int(value: Any, default: int = 0) -> int:
    try:
        return int(float(str(value).replace(",", "")))
    except Exception:
        return default


def parse_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(str(value).replace(",", ""))
    except Exception:
        return default


def tokenize(text: str) -> list[str]:
    tokens = [token for token in TOKEN_RE.findall(clean(text)) if token not in STOPWORDS]
    return [token for token in tokens if not token.isdigit()]


def top_keywords(text: str, limit: int = 8) -> list[str]:
    counter = Counter(tokenize(text))
    return [token for token, _ in counter.most_common(limit)]


def score_rules(text: str, rules: dict[str, list[str]], limit: int) -> list[str]:
    lowered = clean(text).lower()
    scored: list[tuple[str, int]] = []
    for tag, keywords in rules.items():
        hits = sum(1 for keyword in keywords if keyword.lower() in lowered)
        if hits:
            scored.append((tag, hits))
    return [tag for tag, _ in sorted(scored, key=lambda item: item[1], reverse=True)[:limit]]


def phrase_for_tags(tags: list[str]) -> str:
    if not tags:
        return "감정 단서가 약한 일상적 정서"
    mapping = {
        "슬픔": "상실감과 쓸쓸함",
        "그리움": "지나간 관계나 시간을 떠올리는 그리움",
        "분노": "강한 긴장과 억울함",
        "불안": "흔들림과 걱정",
        "무기력": "지친 마음과 낮은 에너지",
        "기쁨": "밝고 활기 있는 기분",
        "설렘": "기대와 두근거림",
        "감사": "따뜻한 고마움",
        "평온": "잔잔하고 안정적인 분위기",
        "희망": "다시 나아가려는 희망",
        "호기심": "새로운 것을 향한 호기심",
    }
    return ", ".join(mapping.get(tag, tag) for tag in tags[:3])


def generation_label(years: set[int]) -> str:
    if not years:
        return "연도 미상"
    median = sorted(years)[len(years) // 2]
    decade = median // 10 * 10
    return f"{decade}년대"


def read_bugs_years(raw_dir: Path) -> dict[str, set[int]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        return {}

    candidates = sorted(
        (path for path in raw_dir.rglob("*.xlsx") if "통합" in path.name),
        key=lambda path: path.stat().st_size,
        reverse=True,
    )
    if not candidates:
        candidates = sorted(raw_dir.rglob("*.xlsx"), key=lambda path: path.stat().st_size, reverse=True)
    if not candidates:
        return {}

    years_by_key: dict[str, set[int]] = defaultdict(set)
    workbook = load_workbook(candidates[0], read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(rows, [])]
        year_idx = headers.index("연도") if "연도" in headers else 0
        title_idx = headers.index("곡") if "곡" in headers else 1
        artist_idx = headers.index("아티스트") if "아티스트" in headers else 2
        for row in rows:
            title = clean(row[title_idx] if len(row) > title_idx else "")
            artist = clean(row[artist_idx] if len(row) > artist_idx else "")
            year = parse_int(row[year_idx] if len(row) > year_idx else "")
            if title and artist and year:
                years_by_key[norm_key(title, artist)].add(year)
    finally:
        workbook.close()
    return years_by_key


def aggregate_melon_chart(path: Path) -> dict[str, dict[str, Any]]:
    records: dict[str, dict[str, Any]] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            title = clean(row.get("title"))
            artist = clean(row.get("artist"))
            if not title or not artist:
                continue
            key = norm_key(title, artist)
            record = records.setdefault(
                key,
                {
                    "title": title,
                    "artist": artist,
                    "albums": Counter(),
                    "genres": Counter(),
                    "years": set(),
                    "song_ids": set(),
                    "appearances": 0,
                    "best_rank": 9999,
                    "max_like_count": 0,
                },
            )
            record["appearances"] += 1
            record["albums"][clean(row.get("album"))] += 1
            record["genres"][clean(row.get("genre_name"))] += 1
            year = parse_int(row.get("year"))
            if year:
                record["years"].add(year)
            if row.get("song_id"):
                record["song_ids"].add(clean(row.get("song_id")))
            record["best_rank"] = min(record["best_rank"], parse_int(row.get("rank"), 9999))
            record["max_like_count"] = max(record["max_like_count"], parse_int(row.get("like_count")))
    return records


def chart_score(record: dict[str, Any]) -> float:
    best_rank = max(1, int(record.get("best_rank") or 9999))
    return (
        float(record.get("appearances") or 0) * 3.0
        + len(record.get("years") or []) * 2.0
        + min(12.0, float(record.get("max_like_count") or 0) / 8000.0)
        + max(0.0, 55.0 - best_rank) / 8.0
    )


def read_lyrics_for_keys(path: Path, target_keys: set[str]) -> dict[str, dict[str, Any]]:
    lyrics_by_key: dict[str, dict[str, Any]] = {}
    if not target_keys:
        return lyrics_by_key
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            title = clean(row.get("title"))
            artist = clean(row.get("artist"))
            if not title or not artist:
                continue
            key = norm_key(title, artist)
            if key not in target_keys or key in lyrics_by_key:
                continue
            raw_lyrics = str(row.get("lyrics") or "")
            processed = preprocess_lyrics(raw_lyrics)
            if not processed:
                continue
            keywords = top_keywords(processed, limit=10)
            emotion_tags = score_rules(processed, EMOTION_RULES, limit=4)
            topic_tags = score_rules(processed, TOPIC_RULES, limit=5)
            lyrics_by_key[key] = {
                "lyrics_hash": hashlib.blake2b(raw_lyrics.encode("utf-8"), digest_size=8).hexdigest(),
                "lyric_keywords": keywords,
                "emotion_tags": emotion_tags or ["평온"],
                "topic_tags": topic_tags or ["일상"],
                "lyric_summary": (
                    f"가사 요약: {phrase_for_tags(emotion_tags)} 중심의 정서가 두드러진다. "
                    f"핵심 키워드: {', '.join(keywords[:7]) or '키워드 부족'}."
                ),
            }
    return lyrics_by_key


def build_music_items(
    chart_path: Path,
    lyrics_path: Path,
    bugs_dir: Path,
    limit: int,
    candidate_multiplier: int = 8,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    records = aggregate_melon_chart(chart_path)
    bugs_years = read_bugs_years(bugs_dir)
    for key, years in bugs_years.items():
        if key in records:
            records[key]["years"].update(years)
            records[key]["source_has_bugs"] = True
    ranked = sorted(records.items(), key=lambda item: chart_score(item[1]), reverse=True)
    target_keys = {key for key, _ in ranked[: max(limit * candidate_multiplier, limit)]}
    lyrics_by_key = read_lyrics_for_keys(lyrics_path, target_keys)

    items: list[dict[str, Any]] = []
    for key, record in ranked:
        lyric = lyrics_by_key.get(key)
        if not lyric:
            continue
        years = set(record.get("years") or [])
        genre = record["genres"].most_common(1)[0][0] if record["genres"] else "가요"
        album = record["albums"].most_common(1)[0][0] if record["albums"] else ""
        generation = generation_label(years)
        source_parts = ["melon_monthly_top50", "melon_song_lyrics"]
        if record.get("source_has_bugs"):
            source_parts.append("bugs_yearly_songs")
        metadata_note = (
            f"{generation} 차트 기반 곡이며, 최고 순위 {record['best_rank']}위, "
            f"차트 등장 {record['appearances']}회로 확인된다."
        )
        summary = (
            f"{record['artist']}의 음악 콘텐츠. {metadata_note} "
            f"{lyric['lyric_summary']} 장르/분류: {genre}."
        )
        topic_tags = list(dict.fromkeys(lyric["topic_tags"] + [genre, generation]))
        emotion_tags = list(dict.fromkeys(lyric["emotion_tags"]))
        tagging_text = "\n".join(
            [
                f"제목: {record['title']}",
                f"아티스트: {record['artist']}",
                "유형: music",
                f"장르: {genre}",
                f"앨범: {album}",
                f"요약: {summary}",
                f"감정 태그: {', '.join(emotion_tags)}",
                f"주제 태그: {', '.join(topic_tags)}",
                f"가사 키워드: {', '.join(lyric['lyric_keywords'])}",
                f"원천: {', '.join(source_parts)}",
            ]
        )
        items.append(
            {
                "content_id": stable_id("music_real", key),
                "content_type": "music",
                "title": record["title"],
                "creator": record["artist"],
                "genre": genre,
                "summary": summary,
                "source": "+".join(source_parts),
                "emotion_tags": emotion_tags,
                "topic_tags": topic_tags,
                "embedding_text": tagging_text,
                "raw_description": summary,
                "processed_description": summary,
                "tagging_text": tagging_text,
                "import_metadata": {
                    "album": album,
                    "years": sorted(years),
                    "generation": generation,
                    "best_rank": record["best_rank"],
                    "appearances": record["appearances"],
                    "max_like_count": record["max_like_count"],
                    "song_ids": sorted(record["song_ids"]),
                    "lyrics_hash": lyric["lyrics_hash"],
                },
            }
        )
        if len(items) >= limit:
            break

    stats = {
        "melon_chart_unique": len(records),
        "bugs_unique": len(bugs_years),
        "lyrics_matched": len(lyrics_by_key),
        "music_items": len(items),
    }
    return items, stats


def infer_book_genre(text: str) -> str:
    checks = [
        ("경제/자기계발", ["부자", "돈", "투자", "경제", "성공", "습관", "자기계발"]),
        ("소설/판타지", ["소설", "마법", "모험", "사랑", "이야기", "주인공", "세계"]),
        ("에세이/위로", ["에세이", "위로", "마음", "삶", "행복", "관계"]),
        ("인문/사회", ["철학", "역사", "사회", "인간", "문명", "정치"]),
        ("과학/지식", ["과학", "우주", "생명", "기술", "지식"]),
        ("학습/교육", ["공부", "학습", "시험", "영어", "교육"]),
        ("건강/생활", ["건강", "요리", "운동", "생활", "몸"]),
    ]
    lowered = text.lower()
    for genre, keywords in checks:
        if any(keyword in lowered for keyword in keywords):
            return genre
    return "베스트셀러"


def aggregate_kyobo(path: Path) -> dict[str, dict[str, Any]]:
    records: dict[str, dict[str, Any]] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            title = clean(row.get("title"))
            author = clean(row.get("author"))
            if not title or not author:
                continue
            key = norm_key(title, author)
            record = records.setdefault(
                key,
                {
                    "title": title,
                    "author": author,
                    "publisher": clean(row.get("publisher")),
                    "years": set(),
                    "best_rank": 9999,
                    "summary": "",
                    "review_keywords": Counter(),
                    "rating": 0.0,
                    "review_count": 0.0,
                    "detail_url": clean(row.get("detail_url")),
                },
            )
            year = parse_int(row.get("year"))
            if year:
                record["years"].add(year)
            record["best_rank"] = min(record["best_rank"], parse_int(row.get("rank"), 9999))
            summary = clean(row.get("summary"))
            if len(summary) > len(record["summary"]):
                record["summary"] = summary
            keyword = clean(row.get("review_keyword"))
            if keyword:
                record["review_keywords"][keyword] += 1
            record["rating"] = max(record["rating"], parse_float(row.get("rating")))
            record["review_count"] = max(record["review_count"], parse_float(row.get("review_count")))
            if not record["detail_url"]:
                record["detail_url"] = clean(row.get("detail_url"))
    return records


def book_score(record: dict[str, Any]) -> float:
    summary_bonus = min(10.0, len(record.get("summary") or "") / 120.0)
    best_rank = max(1, int(record.get("best_rank") or 9999))
    return (
        len(record.get("years") or []) * 4.0
        + max(0.0, 110.0 - best_rank) / 8.0
        + min(8.0, float(record.get("review_count") or 0) / 120.0)
        + float(record.get("rating") or 0) / 2.0
        + summary_bonus
    )


def build_book_items(path: Path, limit: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    records = aggregate_kyobo(path)
    ranked = sorted(records.items(), key=lambda item: book_score(item[1]), reverse=True)
    items: list[dict[str, Any]] = []
    for key, record in ranked:
        base_summary = record["summary"]
        review_keywords = [keyword for keyword, _ in record["review_keywords"].most_common(3)]
        descriptor_source = base_summary or (
            f"{record['title']}은 교보문고 연도별 베스트셀러에 오른 도서다. "
            f"리뷰 키워드: {', '.join(review_keywords) or '없음'}."
        )
        genre = infer_book_genre(" ".join([record["title"], descriptor_source, " ".join(review_keywords)]))
        emotion_tags = score_rules(descriptor_source, EMOTION_RULES, limit=4) or ["호기심"]
        topic_tags = score_rules(descriptor_source + " " + genre, TOPIC_RULES, limit=5) or [genre]
        keywords = top_keywords(descriptor_source, limit=8)
        years = set(record.get("years") or [])
        generation = generation_label(years)
        summary = (
            f"{record['author']}의 도서. {generation} 교보문고 연도별 베스트셀러에 포함되었고 "
            f"최고 순위 {record['best_rank']}위로 확인된다. "
            f"{descriptor_source} 핵심 키워드: {', '.join(keywords) or ', '.join(review_keywords) or '키워드 부족'}."
        )
        topic_tags = list(dict.fromkeys(topic_tags + [genre, generation]))
        tagging_text = "\n".join(
            [
                f"제목: {record['title']}",
                f"저자: {record['author']}",
                "유형: book",
                f"장르: {genre}",
                f"출판사: {record['publisher']}",
                f"요약: {summary}",
                f"감정 태그: {', '.join(emotion_tags)}",
                f"주제 태그: {', '.join(topic_tags)}",
                f"리뷰 키워드: {', '.join(review_keywords)}",
                "원천: kyobo_annual_bestseller",
            ]
        )
        items.append(
            {
                "content_id": stable_id("book_real", key),
                "content_type": "book",
                "title": record["title"],
                "creator": record["author"],
                "genre": genre,
                "summary": summary,
                "source": "kyobo_annual_bestseller",
                "emotion_tags": emotion_tags,
                "topic_tags": topic_tags,
                "embedding_text": tagging_text,
                "raw_description": summary,
                "processed_description": summary,
                "tagging_text": tagging_text,
                "import_metadata": {
                    "publisher": record["publisher"],
                    "years": sorted(years),
                    "generation": generation,
                    "best_rank": record["best_rank"],
                    "rating": record["rating"],
                    "review_count": record["review_count"],
                    "review_keywords": review_keywords,
                    "detail_url": record["detail_url"],
                },
            }
        )
        if len(items) >= limit:
            break
    stats = {
        "kyobo_unique": len(records),
        "book_items": len(items),
        "book_items_with_summary": sum(1 for item in items if len(item["summary"]) >= 80),
    }
    return items, stats


def parse_postgres_values(line: str) -> list[str | None]:
    marker = " VALUES ("
    start = line.find(marker)
    if start < 0:
        return []
    text = line[start + len(marker) :]
    if text.endswith(";\n"):
        text = text[:-2]
    elif text.endswith(";"):
        text = text[:-1]
    if text.endswith(")"):
        text = text[:-1]

    values: list[str | None] = []
    buffer: list[str] = []
    in_string = False
    index = 0
    while index < len(text):
        char = text[index]
        if in_string:
            if char == "'":
                if index + 1 < len(text) and text[index + 1] == "'":
                    buffer.append("'")
                    index += 2
                    continue
                in_string = False
            else:
                buffer.append(char)
        else:
            if char == "'":
                in_string = True
            elif char == ",":
                token = "".join(buffer).strip()
                values.append(None if token.upper() == "NULL" else token)
                buffer = []
            else:
                buffer.append(char)
        index += 1
    token = "".join(buffer).strip()
    values.append(None if token.upper() == "NULL" else token)
    return values


def parse_json_people(raw: str | None, key: str = "peopleNm", limit: int = 3) -> list[str]:
    if not raw:
        return []
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return []
    names: list[str] = []
    if isinstance(parsed, list):
        for item in parsed:
            if not isinstance(item, dict):
                continue
            name = clean(item.get(key))
            if name and name not in names:
                names.append(name)
            if len(names) >= limit:
                break
    return names


def split_genres(genre: str) -> list[str]:
    parts = re.split(r"[,/]", clean(genre))
    return [part.strip() for part in parts if part.strip() and part.strip() != "기타"]


def generation_from_year(year: int) -> str:
    if year <= 0:
        return "연도미상"
    decade = (year // 10) * 10
    return f"{decade}년대"


def infer_movie_profile(title: str, genre: str, nation: str, year: int) -> dict[str, Any]:
    text = " ".join([title, genre, nation]).lower()
    emotion_tags: list[str] = []
    topic_tags: list[str] = ["영화"]
    cue_phrases: list[str] = []

    rules = [
        (
            ["코미디", "comedy"],
            ["기쁨", "즐거움"],
            ["일상", "관계"],
            "유쾌함과 기분 전환의 결이 있어 가볍게 환기하고 싶을 때 어울린다.",
        ),
        (
            ["멜로", "로맨스", "romance"],
            ["설렘", "따뜻함"],
            ["사랑", "관계"],
            "관계와 사랑의 정서를 중심으로 설렘이나 따뜻한 여운을 줄 수 있다.",
        ),
        (
            ["드라마"],
            ["잔잔함", "평범함"],
            ["관계", "일상", "성장"],
            "인물의 관계와 삶을 따라가며 감정을 차분히 정리하기 좋다.",
        ),
        (
            ["애니메이션", "가족"],
            ["기쁨", "따뜻함"],
            ["가족", "성장"],
            "따뜻함과 성장의 분위기가 있어 부담 낮은 감상에 적합하다.",
        ),
        (
            ["액션", "범죄", "스릴러", "느와르", "미스터리"],
            ["호기심"],
            ["갈등", "도전"],
            "긴장감과 갈등의 흐름이 있어 집중해서 몰입하고 싶을 때 맞다.",
        ),
        (
            ["공포", "호러"],
            ["불안", "호기심"],
            ["갈등"],
            "어두운 긴장과 불안의 분위기가 강해 자극적인 감상을 원할 때 적합하다.",
        ),
        (
            ["sf", "s/f", "판타지"],
            ["호기심", "경외", "희망"],
            ["과학", "미래", "우주"],
            "상상력과 미지의 세계를 다루며 호기심과 몰입감을 키운다.",
        ),
        (
            ["다큐멘터리"],
            ["호기심", "경외"],
            ["학습", "사회"],
            "정보와 현실 맥락을 따라가며 새롭게 알아가는 감각이 강하다.",
        ),
        (
            ["사극", "전쟁", "역사"],
            ["잔잔함", "호기심"],
            ["기억", "사회", "갈등"],
            "역사적 기억과 갈등을 돌아보게 하는 성찰형 감상에 가깝다.",
        ),
        (
            ["뮤지컬", "공연"],
            ["즐거움", "따뜻함"],
            ["예술", "음악"],
            "음악과 무대성이 있어 감정을 밝게 움직이는 감상에 어울린다.",
        ),
    ]
    for keywords, emotions, topics, phrase in rules:
        if any(keyword in text for keyword in keywords):
            emotion_tags.extend(emotions)
            topic_tags.extend(topics)
            cue_phrases.append(phrase)

    if not emotion_tags:
        emotion_tags = ["호기심", "평범함"]
        topic_tags.extend(["예술"])
        cue_phrases.append("장르와 제작 정보를 바탕으로 새로운 감상 대상을 탐색하기 좋다.")

    for genre_part in split_genres(genre):
        topic_tags.append(genre_part)
    if nation:
        topic_tags.append(nation)
    if year:
        topic_tags.append(generation_from_year(year))

    return {
        "emotion_tags": list(dict.fromkeys(emotion_tags))[:4],
        "topic_tags": list(dict.fromkeys(topic_tags))[:8],
        "cue": " ".join(cue_phrases),
    }


def parse_kobis_movie_line(line: str) -> dict[str, Any] | None:
    if not line.startswith("INSERT INTO public.kobis_movie VALUES"):
        return None
    values = parse_postgres_values(line)
    if len(values) < 22:
        return None
    movie_cd = clean(values[0])
    title = clean(values[1])
    if not movie_cd or not title:
        return None

    year = parse_int(values[3])
    open_date = clean(values[4])
    type_name = clean(values[5])
    production_status = clean(values[6])
    nation = clean(values[9] or values[7])
    genre = clean(values[8] or values[10])
    original_title = clean(values[14])
    runtime = parse_int(values[15])
    directors = parse_json_people(values[11], limit=3)
    actors = parse_json_people(values[18], limit=4)
    if not genre or genre == "기타":
        return None

    profile = infer_movie_profile(title, genre, nation, year)
    director_text = ", ".join(directors) if directors else "감독 정보 미상"
    actor_text = ", ".join(actors) if actors else "주요 출연 정보 미상"
    generation = generation_from_year(year)
    metadata_parts = [
        f"{year}년" if year else "제작연도 미상",
        nation or "국가 미상",
        genre,
    ]
    if type_name:
        metadata_parts.append(type_name)
    if production_status:
        metadata_parts.append(production_status)
    if runtime:
        metadata_parts.append(f"{runtime}분")
    if original_title:
        metadata_parts.append(f"원제/별칭: {original_title}")

    summary = (
        f"KOBIS 영화 메타데이터 기반 콘텐츠. {' · '.join(metadata_parts)}. "
        f"감독: {director_text}. 출연/참여: {actor_text}. "
        f"추천 태깅 힌트: {profile['cue']}"
    )
    tagging_text = "\n".join(
        [
            f"제목: {title}",
            "유형: movie",
            f"장르: {genre}",
            f"국가: {nation}",
            f"제작연도: {year or '미상'}",
            f"세대: {generation}",
            f"감독: {director_text}",
            f"출연: {actor_text}",
            f"요약: {summary}",
            f"감정 태그 후보: {', '.join(profile['emotion_tags'])}",
            f"주제 태그 후보: {', '.join(profile['topic_tags'])}",
            "원천: kobis_seed_sql",
        ]
    )
    return {
        "content_id": f"movie_kobis_{movie_cd}",
        "content_type": "movie",
        "title": title,
        "creator": director_text,
        "genre": genre.replace(",", "/"),
        "summary": summary,
        "source": "kobis_seed_sql",
        "emotion_tags": profile["emotion_tags"],
        "topic_tags": profile["topic_tags"],
        "embedding_text": tagging_text,
        "raw_description": summary,
        "processed_description": summary,
        "tagging_text": tagging_text,
        "import_metadata": {
            "movie_cd": movie_cd,
            "movie_nm_en": clean(values[2]),
            "open_date": open_date,
            "type_name": type_name,
            "production_status": production_status,
            "nation": nation,
            "year": year,
            "runtime": runtime,
            "directors": directors,
            "actors": actors,
        },
    }


def movie_score(item: dict[str, Any]) -> float:
    metadata = item.get("import_metadata") or {}
    year = int(metadata.get("year") or 0)
    genre = item.get("genre") or ""
    nation = metadata.get("nation") or ""
    score = 0.0
    if year:
        score += max(0, year - 1980) / 2.5
    if metadata.get("production_status") == "개봉":
        score += 22.0
    if nation == "한국":
        score += 18.0
    if metadata.get("type_name") == "장편":
        score += 8.0
    if metadata.get("directors"):
        score += 8.0
    if metadata.get("actors"):
        score += 5.0
    if metadata.get("runtime"):
        score += 3.0
    if genre and genre != "기타":
        score += 12.0
    if any(keyword in genre for keyword in ["드라마", "코미디", "멜로", "로맨스", "액션", "스릴러", "SF", "애니메이션"]):
        score += 5.0
    return score


def build_movie_items(path: Path, limit: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if limit <= 0 or not path.exists():
        return [], {
            "kobis_seed_exists": path.exists(),
            "kobis_seed_path": str(path),
            "kobis_movie_rows": 0,
            "movie_items": 0,
        }

    candidates: list[dict[str, Any]] = []
    movie_rows = 0
    with path.open("r", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            if not line.startswith("INSERT INTO public.kobis_movie VALUES"):
                continue
            movie_rows += 1
            item = parse_kobis_movie_line(line)
            if item:
                candidates.append(item)

    ranked = sorted(candidates, key=movie_score, reverse=True)
    items = ranked[:limit]
    stats = {
        "kobis_seed_exists": True,
        "kobis_seed_path": str(path),
        "kobis_movie_rows": movie_rows,
        "kobis_movie_candidates": len(candidates),
        "movie_items": len(items),
        "movie_items_korean": sum(
            1 for item in items if (item.get("import_metadata") or {}).get("nation") == "한국"
        ),
    }
    return items, stats


def write_processed(items: list[dict[str, Any]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for item in items:
            safe_item = dict(item)
            safe_item.pop("import_metadata", None)
            handle.write(json.dumps(safe_item, ensure_ascii=False) + "\n")


def filter_existing_items(database: Database, items: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int]:
    if not items:
        return [], 0
    ids = [str(item["content_id"]) for item in items]
    existing: set[str] = set()
    with database.connect() as connection:
        for start in range(0, len(ids), 900):
            chunk = ids[start : start + 900]
            placeholders = ",".join("?" for _ in chunk)
            rows = connection.execute(
                f"SELECT content_id FROM content_items WHERE content_id IN ({placeholders})",
                chunk,
            ).fetchall()
            existing.update(str(row[0]) for row in rows)
    return [item for item in items if str(item["content_id"]) not in existing], len(existing)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Import real music/book/movie culture data into the MVP DB.")
    parser.add_argument("--music-limit", type=int, default=500)
    parser.add_argument("--book-limit", type=int, default=500)
    parser.add_argument("--movie-limit", type=int, default=300)
    parser.add_argument("--skip-music", action="store_true")
    parser.add_argument("--skip-books", action="store_true")
    parser.add_argument("--skip-movies", action="store_true")
    parser.add_argument("--movie-seed-path", type=Path, default=KOBIS_SEED_PATH)
    parser.add_argument("--skip-db", action="store_true")
    parser.add_argument("--skip-existing", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--processed-output", type=Path, default=PROCESSED_DIR / "culture_import_music_book_movie.jsonl")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    items: list[dict[str, Any]] = []
    stats: dict[str, Any] = {}

    if not args.skip_music:
        music_items, music_stats = build_music_items(
            MELON_CHART_PATH,
            MELON_LYRICS_PATH,
            RAW_DIR / "bugs_연도별가요",
            max(0, args.music_limit),
        )
        items.extend(music_items)
        stats.update(music_stats)

    if not args.skip_books:
        book_items, book_stats = build_book_items(KYOBO_PATH, max(0, args.book_limit))
        items.extend(book_items)
        stats.update(book_stats)

    if not args.skip_movies:
        movie_items, movie_stats = build_movie_items(args.movie_seed_path, max(0, args.movie_limit))
        items.extend(movie_items)
        stats.update(movie_stats)

    existing_skipped = 0
    database: Database | None = None
    if args.skip_existing:
        database = Database()
        database.initialize()
        items, existing_skipped = filter_existing_items(database, items)

    write_processed(items, args.processed_output)
    db_result = {"skipped": bool(args.skip_db or args.dry_run)}
    if not args.skip_db and not args.dry_run:
        if database is None:
            database = Database()
            database.initialize()
        db_result = database.upsert_content_items(items, reset_tagging=True)

    report = {
        "music_limit": args.music_limit,
        "book_limit": args.book_limit,
        "movie_limit": args.movie_limit,
        "processed_output": str(args.processed_output),
        "total_items": len(items),
        "content_type_counts": dict(Counter(item["content_type"] for item in items)),
        "stats": stats,
        "existing_skipped": existing_skipped,
        "db": db_result,
        "raw_lyrics_persisted": False,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
